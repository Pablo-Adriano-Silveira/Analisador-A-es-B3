"""Exportação para CSV (padrão Excel brasileiro: ';' e vírgula decimal, UTF-8 com BOM)
e para Excel (.xlsx, com números reais e não texto)."""
from __future__ import annotations

import csv
import re
from typing import Any, Dict, List, Sequence, Tuple

from .models import DataError
from .tables import Column

XLSX_FORMATS = {
    "int": "#,##0", "vol": "#,##0", "pct": '0.00"%"', "score": "0.00", "money": '"R$" #,##0.00',
    "pp": '+0.00" p.p.";-0.00" p.p.";0.00" p.p."',
}


def _csv_cell(kind: str, v: Any) -> str:
    if v is None:
        return ""
    if kind == "bool":
        return "Sim" if v else "Não"
    if kind in ("int", "vol"):
        return str(int(round(v)))
    if kind == "money":
        return f"{v:.2f}".replace(".", ",")
    if kind in ("pct", "score", "pp"):
        return f"{v:.4f}".replace(".", ",")
    return str(v)


def export_csv(path: str, columns: Sequence[Column], rows: List[Dict[str, Any]]) -> None:
    try:
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow([c.title for c in columns])
            for r in rows:
                w.writerow([_csv_cell(c.kind, r.get(c.key)) for c in columns])
    except OSError as exc:
        raise DataError(f"Não foi possível gravar o arquivo:\n{path}\n({exc.strerror or exc})") from exc


def _sheet_name(name: str, used: set) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "-", name)[:31] or "Planilha"
    out, n = base, 2
    while out in used:
        out = f"{base[:28]}_{n}"
        n += 1
    used.add(out)
    return out


def export_xlsx(path: str,
                sheets: Sequence[Tuple[str, Sequence[Column], List[Dict[str, Any]]]],
                summary: Sequence[Tuple[str, Any]] = ()) -> None:
    """sheets: [(nome, colunas, linhas)]. summary: pares (rótulo, valor) para a aba 'Resumo'."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    head_fill = PatternFill("solid", fgColor="0F2742")
    head_font = Font(bold=True, color="FFFFFF")

    for name, columns, rows in sheets:
        ws = wb.create_sheet(_sheet_name(name, used))
        ws.append([c.title for c in columns])
        for cell in ws[1]:
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r in rows:
            vals = []
            for c in columns:
                v = r.get(c.key)
                if c.kind == "bool" and v is not None:
                    v = "Sim" if v else "Não"
                vals.append(v)
            ws.append(vals)
        for j, c in enumerate(columns, start=1):
            fmt = XLSX_FORMATS.get(c.kind)
            if fmt:
                for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    row[0].number_format = fmt
            ws.column_dimensions[get_column_letter(j)].width = max(11, min(30, len(c.title) + 4))
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 32

    if summary:
        ws = wb.create_sheet(_sheet_name("Resumo", used))
        ws.append(["Item", "Valor"])
        for cell in ws[1]:
            cell.fill, cell.font = head_fill, head_font
        for k, v in summary:
            ws.append([k, v])
        ws.column_dimensions["A"].width = 52
        ws.column_dimensions["B"].width = 26
        ws.append([])
        ws.append(["Aviso: ferramenta de análise estatística sobre dados históricos; "
                   "não é recomendação de investimento nem garantia de desempenho futuro."])

    if not wb.sheetnames:
        raise DataError("Não há dados para exportar.")
    try:
        wb.save(path)
    except OSError as exc:
        raise DataError(f"Não foi possível gravar o arquivo:\n{path}\n({exc.strerror or exc})") from exc

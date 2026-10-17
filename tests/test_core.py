"""Testes do núcleo com dados sintéticos. Execução: python -m unittest discover -s tests -v"""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from analisador_acoes.core.comparison import compare, sort_matches, SORT_OPTIONS  # noqa: E402
from analisador_acoes.core.csv_loader import (load_csv, parse_count, parse_decimal,  # noqa: E402
                                              parse_volume, normalize_header)
from analisador_acoes.core.exporter import export_csv, export_xlsx  # noqa: E402
from analisador_acoes.core.models import DataError  # noqa: E402
from analisador_acoes.core.ranking import ScoringConfig, rank_records, wilson_lower  # noqa: E402
from analisador_acoes.core.tables import (RANKING_COLUMNS, ranking_rows,  # noqa: E402
                                          comparison_columns, comparison_rows)

HEADER = ",Ativo,Qtd Gain,Taxa de Acerto,Resultado Acumulado,Gain Médio (R$),Gain Médio (%),Qtd Loss,Maior Loss (R$),Maior Loss (%),Méd. Volume Leilão"


def line(i, t, g, l, vol, extra=True):
    taxa = (100.0 * g / (g + l)) if (g + l) else 0
    return f'{i},{t},{g},{taxa},1.0,"R$ 0,10",1.0,{l},"R$ -0,10",-1.0,{vol}'


def write(content, enc="utf-8-sig", name="t.csv"):
    d = tempfile.mkdtemp()
    p = os.path.join(d, name)
    with open(p, "w", encoding=enc, newline="") as fh:
        fh.write(content)
    return p


class ParseTests(unittest.TestCase):
    def test_volume(self):
        self.assertEqual(parse_volume("29.869.975"), 29869975)
        self.assertEqual(parse_volume("58.448"), 58448)       # milhar, não 58,448
        self.assertEqual(parse_volume("314513"), 314513)
        self.assertEqual(parse_volume("1.234,5"), 1234.5)
        self.assertEqual(parse_volume("0"), 0)
        self.assertIsNone(parse_volume(""))
        self.assertIsNone(parse_volume("abc"))
        self.assertIsNone(parse_volume("-5"))

    def test_decimal_and_count(self):
        self.assertAlmostEqual(parse_decimal("92.857"), 92.857)
        self.assertAlmostEqual(parse_decimal("92,857"), 92.857)
        self.assertAlmostEqual(parse_decimal("1.234,50"), 1234.5)
        self.assertAlmostEqual(parse_decimal("85%"), 85)
        self.assertIsNone(parse_decimal("nan"))
        self.assertEqual(parse_count("13"), 13)
        self.assertEqual(parse_count("13.0"), 13)
        self.assertIsNone(parse_count("-1"))
        self.assertIsNone(parse_count("2.5"))
        self.assertIsNone(parse_count(""))

    def test_header(self):
        self.assertEqual(normalize_header("Méd. Volume Leilão"), "med volume leilao")


class LoaderTests(unittest.TestCase):
    def test_basic_bom_and_extra_columns(self):
        p = write("\n".join([HEADER, line(0, "ABEV3", 4, 0, "29.869.975"), line(1, "PETR4", 8, 2, "5.000.000")]))
        r = load_csv(p)
        self.assertEqual([x.ticker for x in r.records], ["ABEV3", "PETR4"])
        self.assertEqual(r.records[0].volume, 29869975)
        self.assertIn("Resultado Acumulado", r.extra_columns)
        self.assertEqual(r.encoding, "utf-8-sig")

    def test_semicolon_cp1252_and_header_variants(self):
        txt = "Ticker;QTD GAIN;Qtd Loss;taxa de acerto;Med. Volume Leilao\npetr4 ;5;5;50;\"1.000.000\"\n"
        r = load_csv(write(txt, enc="cp1252"))
        self.assertEqual(r.delimiter, ";")
        self.assertEqual(r.records[0].ticker, "PETR4")      # espaços e minúsculas normalizados

    def test_missing_columns_message(self):
        p = write("Ativo,Qtd Gain\nABEV3,4\n")
        with self.assertRaises(DataError) as cm:
            load_csv(p)
        msg = str(cm.exception)
        self.assertIn("Qtd Loss", msg)
        self.assertIn("Méd. Volume Leilão", msg)

    def test_empty_file(self):
        with self.assertRaises(DataError):
            load_csv(write(""))
        with self.assertRaises(DataError):
            load_csv(write(HEADER + "\n"))          # só cabeçalho

    def test_missing_file(self):
        with self.assertRaises(DataError):
            load_csv("/nao/existe.csv")

    def test_missing_values_not_zero_filled(self):
        p = write("\n".join([HEADER, line(0, "AAAA3", 5, 5, ""), line(1, "BBBB3", 5, 5, "1.000.000")]))
        r = load_csv(p)
        self.assertEqual([x.ticker for x in r.records], ["BBBB3"])
        self.assertEqual(len(r.invalid_rows), 1)
        self.assertIn("volume", r.invalid_rows[0][2].lower())

    def test_duplicates_keep_most_ops_and_warn(self):
        p = write("\n".join([HEADER, line(0, "AAAA3", 2, 2, "1.000.000"),
                             line(1, "AAAA3", 10, 5, "1.000.000"), line(2, "BBBB3", 6, 4, "1.000.000")]))
        r = load_csv(p)
        self.assertEqual(len(r.records), 2)
        self.assertEqual(next(x for x in r.records if x.ticker == "AAAA3").ops, 15)
        self.assertEqual(r.duplicates, [("AAAA3", 2)])
        self.assertTrue(any("duplicados" in w for w in r.warnings))

    def test_does_not_modify_original(self):
        p = write("\n".join([HEADER, line(0, "ABEV3", 4, 0, "1.000.000")]))
        with open(p, "rb") as fh:
            before = fh.read()
        load_csv(p)
        with open(p, "rb") as fh:
            self.assertEqual(before, fh.read())


def rec_file(rows):
    return load_csv(write("\n".join([HEADER] + [line(i, *r) for i, r in enumerate(rows)])))


class RankingTests(unittest.TestCase):
    def test_volume_boundary_is_strict(self):
        r = rec_file([("AAAA3", 8, 2, "80.000"), ("BBBB3", 8, 2, "80.001"), ("CCCC3", 8, 2, "79.999")])
        res = rank_records(r.records, ScoringConfig())
        self.assertEqual([x.record.ticker for x in res.ranked], ["BBBB3"])
        self.assertEqual(len(res.excluded_volume), 2)
        self.assertEqual(res.volume_ok, 1)

    def test_min_ops_and_zero_gain_filters(self):
        r = rec_file([("AAAA3", 4, 0, "1.000.000"), ("BBBB3", 0, 12, "1.000.000"), ("CCCC3", 9, 1, "1.000.000")])
        res = rank_records(r.records, ScoringConfig())
        self.assertEqual([x.record.ticker for x in res.ranked], ["CCCC3"])
        reasons = {e.record.ticker: e.reason for e in res.excluded_filters}
        self.assertIn("Operações", reasons["AAAA3"])
        self.assertIn("Gain", reasons["BBBB3"])

    def test_small_sample_does_not_beat_large_sample(self):
        # 10/10 (100%) x 17/20 (85%): Wilson favorece a amostra maior? 10/10 -> 0.72; 17/20 -> 0.64
        self.assertGreater(wilson_lower(10, 10), wilson_lower(17, 20))
        # mas 4/4 (100%) é MENOR que 17/20 — a amostra pequena é penalizada
        self.assertLess(wilson_lower(4, 4), wilson_lower(17, 20))

    def test_never_pads_below_top_n(self):
        r = rec_file([(f"AA{i:02d}3", 7, 3, "1.000.000") for i in range(5)] + [("ZZZZ3", 9, 1, "10")])
        res = rank_records(r.records, ScoringConfig(top_n=40))
        self.assertEqual(len(res.rows), 5)
        self.assertNotIn("ZZZZ3", [x.record.ticker for x in res.rows])

    def test_top_n_limit_and_order(self):
        rows = [(f"T{i:03d}3", 10 + i, 10 - (i % 5), "1.000.000") for i in range(60)]
        res = rank_records(rec_file(rows).records, ScoringConfig(top_n=40))
        self.assertEqual(len(res.rows), 40)
        scores = [x.score for x in res.rows]
        self.assertEqual(scores, sorted(scores, reverse=True))
        self.assertEqual([x.pos for x in res.rows], list(range(1, 41)))
        self.assertTrue(all(0 <= s <= 100 for s in scores))

    def test_score_formula_matches_manual_calculation(self):
        # 3 ações: calculo manual dos percentis
        r = rec_file([("AAAA3", 9, 1, "1.000.000"), ("BBBB3", 6, 4, "1.000.000"), ("CCCC3", 8, 8, "1.000.000")])
        res = rank_records(r.records, ScoringConfig())
        a = next(x for x in res.ranked if x.record.ticker == "AAAA3")
        # AAAA3: melhor Wilson (3/3), melhor Qtd Gain (3/3), menos Qtd Loss (3/3) -> nota 100
        self.assertAlmostEqual(a.score, 100.0)
        c = next(x for x in res.ranked if x.record.ticker == "CCCC3")
        # CCCC3: pior Wilson (1/3), Gain intermediário (2/3), pior Loss (1/3)
        self.assertAlmostEqual(c.score, 100 * (0.7 * 1 / 3 + 0.2 * 2 / 3 + 0.1 * 1 / 3))

    def test_empty_eligible(self):
        res = rank_records(rec_file([("AAAA3", 1, 1, "1")]).records, ScoringConfig())
        self.assertEqual(res.rows, [])
        self.assertEqual(res.eligible, 0)


class CompareTests(unittest.TestCase):
    def setUp(self):
        self.a = rec_file([("AAAA3", 9, 1, "1.000.000"), ("BBBB3", 6, 4, "2.000.000"),
                           ("CCCC3", 7, 3, "500.000"), ("ONLYA3", 8, 2, "1.000.000"),
                           ("LOWV3", 8, 2, "50.000"), ("FEWO3", 3, 1, "1.000.000")])
        self.b = rec_file([("BBBB3", 12, 8, "2.000.000"), ("AAAA3", 15, 2, "1.000.000"),
                           ("CCCC3", 9, 6, "70.000"), ("ONLYB3", 8, 2, "1.000.000"),
                           ("LOWV3", 8, 2, "50.000"), ("FEWO3", 8, 4, "1.000.000")])

    def test_matches_by_ticker_regardless_of_position(self):
        c = compare(self.a, self.b, ScoringConfig())
        self.assertEqual({m.ticker for m in c.matches}, {"AAAA3", "BBBB3"})
        self.assertEqual(c.n_both, 5)
        self.assertEqual(c.only_a, ["ONLYA3"])
        self.assertEqual(c.only_b, ["ONLYB3"])

    def test_volume_rule_requires_both_files(self):
        c = compare(self.a, self.b, ScoringConfig())
        self.assertEqual(c.dropped_volume, 2)             # LOWV3 e CCCC3 (70.000 em B)
        self.assertEqual(c.volume_divergences, 1)         # CCCC3 difere entre os arquivos
        self.assertNotIn("CCCC3", {m.ticker for m in c.matches})
        self.assertTrue(any("volume diferente" in n for n in c.notes))

    def test_min_ops_in_either_file_drops_match(self):
        c = compare(self.a, self.b, ScoringConfig())
        self.assertEqual(c.dropped_filters, 1)            # FEWO3 tem só 4 operações em A
        self.assertNotIn("FEWO3", {m.ticker for m in c.matches})

    def test_limit_top_n(self):
        c = compare(self.a, self.b, ScoringConfig(), limit_top_n=1)
        self.assertLessEqual(len(c.matches), 1)
        self.assertGreaterEqual(c.dropped_top_n, 1)

    def test_sorting_options_and_columns(self):
        c = compare(self.a, self.b, ScoringConfig())
        for opt in SORT_OPTIONS:
            self.assertEqual(len(sort_matches(c.matches, opt)), len(c.matches))
        rows = comparison_rows(c.matches)
        keys = {col.key for col in comparison_columns("2%", "1%")}
        self.assertTrue(set(rows[0]).issubset(keys))
        by_mean = sort_matches(c.matches, "Nota média")
        self.assertGreaterEqual(by_mean[0].score_mean, by_mean[-1].score_mean)

    def test_disjoint_files_give_zero_matches(self):
        x = rec_file([("AAAA3", 9, 1, "1.000.000")])
        y = rec_file([("BBBB3", 9, 1, "1.000.000")])
        c = compare(x, y, ScoringConfig())
        self.assertEqual((c.n_both, len(c.matches)), (0, 0))


class ExportTests(unittest.TestCase):
    def test_exports_roundtrip(self):
        r = rec_file([(f"T{i:03d}3", 10 + i, 5, "1.000.000") for i in range(12)])
        res = rank_records(r.records, ScoringConfig())
        rows = ranking_rows(res)
        d = tempfile.mkdtemp()
        pc, px = os.path.join(d, "r.csv"), os.path.join(d, "r.xlsx")
        export_csv(pc, RANKING_COLUMNS, rows)
        export_xlsx(px, [("Compra", RANKING_COLUMNS, rows)], [("Total", 12)])
        with open(pc, encoding="utf-8-sig") as fh:
            txt = fh.read().splitlines()
        self.assertEqual(len(txt), 13)
        self.assertTrue(txt[0].startswith("Pos.;Ativo;"))
        from openpyxl import load_workbook
        wb = load_workbook(px)
        ws = wb["Compra"]
        self.assertEqual(ws.max_row, 13)
        self.assertEqual(ws["B2"].value, rows[0]["ticker"])
        self.assertIsInstance(ws["I2"].value, float)       # número real, não texto
        self.assertIn("Resumo", wb.sheetnames)


if __name__ == "__main__":
    unittest.main()

"""Testes do filtro de valor da ação e do serviço brapi (sem internet: respostas simuladas)."""
import json
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from analisador_acoes.core.analysis import gather_prices, resolve_token  # noqa: E402
from analisador_acoes.core.comparison import compare  # noqa: E402
from analisador_acoes.core.csv_loader import load_csv  # noqa: E402
from analisador_acoes.core.exporter import export_csv, export_xlsx  # noqa: E402
from analisador_acoes.core.formatting import fmt_value  # noqa: E402
from analisador_acoes.core.models import DataError  # noqa: E402
from analisador_acoes.core.price_service import (AuthError, BrapiClient, NetworkError, PriceCache,  # noqa: E402
                                                 QuotaError, fetch_prices, parse_price)
from analisador_acoes.core.ranking import ScoringConfig, candidates, rank_records  # noqa: E402
from analisador_acoes.core.settings import load_settings, save_settings  # noqa: E402
from analisador_acoes.core.tables import (RANKING_COLUMNS, comparison_columns, comparison_rows,  # noqa: E402
                                          comparison_summary, excluded_rows, ranking_rows)
from tests.test_core import HEADER, line, write  # noqa: E402


def body(symbol="PETR4", price=38.5, wrap=True):
    data = {"symbol": symbol, "shortName": "X", "regularMarketPrice": price, "currency": "BRL"}
    return json.dumps({"results": [{"data": data} if wrap else data]}).encode()


def rec_file(rows):
    return load_csv(write("\n".join([HEADER] + [line(i, *r) for i, r in enumerate(rows)])))


class ParseTests(unittest.TestCase):
    def test_parse_price(self):
        self.assertEqual(parse_price(body(), "PETR4"), 38.5)
        self.assertEqual(parse_price(body(wrap=False), "PETR4"), 38.5)          # formato sem "data"
        self.assertIsNone(parse_price(body(price=None), "PETR4"))              # ausente -> None (não zero)
        self.assertIsNone(parse_price(body(price=0), "PETR4"))
        self.assertIsNone(parse_price(body(price=-1), "PETR4"))
        self.assertIsNone(parse_price(body(symbol="VALE3"), "PETR4"))          # ticker diferente
        self.assertIsNone(parse_price(b"not json", "PETR4"))
        self.assertIsNone(parse_price(b'{"results": []}', "PETR4"))
        self.assertIsNone(parse_price(b'{"results": [{"data": {"regularMarketPrice": "abc"}}]}', "PETR4"))


class ClientTests(unittest.TestCase):
    def client(self, status, payload=b"", calls=None):
        def fake(url, token, timeout):
            if calls is not None:
                calls.append((url, token))
            return status, payload
        return BrapiClient("abc123", fetch=fake, sleep=lambda s: None)

    def test_ok_url_and_token(self):
        calls = []
        self.assertEqual(self.client(200, body(), calls).fetch("PETR4"), 38.5)
        self.assertEqual(calls[0][0], "https://brapi.dev/api/v2/stocks/quote?symbols=PETR4")
        self.assertEqual(calls[0][1], "abc123")

    def test_errors(self):
        with self.assertRaises(AuthError):
            self.client(401).fetch("PETR4")
        with self.assertRaises(AuthError):
            self.client(403).fetch("PETR4")
        with self.assertRaises(QuotaError):
            self.client(429).fetch("PETR4")
        self.assertIsNone(self.client(404).fetch("XXXX3"))
        self.assertIsNone(self.client(400).fetch("XXXX3"))
        with self.assertRaises(AuthError):
            BrapiClient("   ")

    def test_retry_on_500_then_ok(self):
        seq = iter([(500, b""), (200, body())])
        c = BrapiClient("k", fetch=lambda u, t, to: next(seq), sleep=lambda s: None)
        self.assertEqual(c.fetch("PETR4"), 38.5)

    def test_429_waits_retry_after_then_succeeds(self):
        seq = iter([(429, b"", {"retry-after": "3"}), (429, b""), (200, body(), {})])
        waits, msgs = [], []
        c = BrapiClient("k", fetch=lambda u, t, to: next(seq), sleep=waits.append, notify=msgs.append)
        self.assertEqual(c.fetch("PETR4"), 38.5)
        self.assertEqual(waits, [3.0, 4.0])               # Retry-After da API; depois espera crescente (2s x 2)
        self.assertEqual(len(msgs), 2)

    def test_429_persistent_raises_after_retries(self):
        calls, waits = [], []
        def always429(u, t, to):
            calls.append(1)
            return 429, b"", {"retry-after": "1"}
        c = BrapiClient("k", fetch=always429, sleep=waits.append)
        with self.assertRaises(QuotaError) as cm:
            c.fetch("PETR4")
        self.assertEqual(len(calls), 5)                   # 1 tentativa + 4 repetições
        self.assertIn("429", str(cm.exception))

    def test_batch_is_sequential_by_default(self):
        import inspect
        self.assertEqual(inspect.signature(fetch_prices).parameters["workers"].default, 1)

    def test_retry_after_clamped(self):
        from analisador_acoes.core.price_service import _retry_after
        self.assertEqual(_retry_after({"retry-after": "999"}, 1), 30.0)
        self.assertEqual(_retry_after({"retry-after": "0"}, 1), 0.5)
        self.assertEqual(_retry_after({}, 3), 6.0)
        self.assertEqual(_retry_after({"retry-after": "abc"}, 2), 4.0)

    def test_network_error_propagates(self):
        def boom(u, t, to):
            raise NetworkError("sem internet")
        with self.assertRaises(NetworkError):
            BrapiClient("k", fetch=boom).fetch("PETR4")


class FakeClient:
    def __init__(self, prices, fail_on=None, exc=None):
        self.prices, self.fail_on, self.exc, self.calls = prices, fail_on, exc, []

    def fetch(self, t):
        self.calls.append(t)
        if t == self.fail_on:
            raise self.exc
        return self.prices.get(t)


class BatchTests(unittest.TestCase):
    def cache(self, ttl=100, clock=None):
        d = tempfile.mkdtemp()
        return PriceCache(os.path.join(d, "c.json"), ttl=ttl, now=clock or (lambda: 1000.0))

    def test_cache_hit_avoids_calls_and_persists(self):
        cache = self.cache()
        fc = FakeClient({"AAAA3": 5.0, "BBBB3": 3.0})
        b = fetch_prices(["AAAA3", "BBBB3"], fc, cache, workers=1)
        self.assertEqual(b.prices, {"AAAA3": 5.0, "BBBB3": 3.0})
        fc2 = FakeClient({})
        b2 = fetch_prices(["AAAA3", "BBBB3"], fc2, PriceCache(cache.path, now=lambda: 1000.0), workers=1)
        self.assertEqual(fc2.calls, [])
        self.assertEqual(b2.from_cache, 2)
        b3 = fetch_prices(["AAAA3"], fc2, PriceCache(cache.path, now=lambda: 1000.0), workers=1, refresh=True)
        self.assertEqual(fc2.calls, ["AAAA3"])                    # refresh ignora o cache
        self.assertEqual(b3.missing, {"AAAA3": "sem cotação na API"})   # e não inventa preço

    def test_cache_expires(self):
        t = [1000.0]
        cache = self.cache(ttl=50, clock=lambda: t[0])
        cache.put("AAAA3", 5.0)
        t[0] = 1049.0
        self.assertEqual(cache.get("AAAA3"), 5.0)
        t[0] = 1051.0
        self.assertIsNone(cache.get("AAAA3"))

    def test_missing_is_not_zero(self):
        b = fetch_prices(["AAAA3", "ZZZZ3"], FakeClient({"AAAA3": 5.0}), self.cache(), workers=1)
        self.assertNotIn("ZZZZ3", b.prices)
        self.assertIn("ZZZZ3", b.missing)
        self.assertIsNone(b.fatal)

    def test_fatal_error_aborts_and_keeps_partial(self):
        tick = ["A1", "A2", "A3", "A4", "A5"]
        fc = FakeClient({t: 5.0 for t in tick}, fail_on="A3", exc=QuotaError("limite"))
        b = fetch_prices(tick, fc, self.cache(), workers=1)
        self.assertEqual(b.fatal, "limite")
        self.assertEqual(set(b.prices), {"A1", "A2"})
        self.assertNotIn("A4", fc.calls)                           # abortou o resto

    def test_gather_prices_only_candidates_and_error(self):
        r = rec_file([("AAAA3", 8, 2, "1.000.000"), ("LOWV3", 8, 2, "50"), ("FEWO3", 2, 1, "1.000.000"),
                      ("BBBB3", 7, 3, "1.000.000")])
        cfg = ScoringConfig(min_price=4.0)
        self.assertEqual({x.ticker for x in candidates(r.records, cfg)}, {"AAAA3", "BBBB3"})
        fc = FakeClient({"AAAA3": 5.0, "BBBB3": 3.0})
        prices = gather_prices([r.records], cfg, "tok", client=fc, cache=self.cache())
        self.assertEqual(sorted(fc.calls), ["AAAA3", "BBBB3"])    # volume baixo e poucas ops não consultam a API
        self.assertEqual(prices, {"AAAA3": 5.0, "BBBB3": 3.0})
        bad = FakeClient({}, fail_on="AAAA3", exc=AuthError("chave inválida"))
        with self.assertRaises(DataError) as cm:
            gather_prices([r.records], cfg, "tok", client=bad, cache=self.cache())
        self.assertIn("chave inválida", str(cm.exception))
        with self.assertRaises(DataError):
            gather_prices([r.records], cfg, "", cache=self.cache())      # sem chave

    def test_resolve_token_env(self):
        os.environ["BRAPI_TOKEN"] = "envtok"
        try:
            self.assertEqual(resolve_token("  "), "envtok")
            self.assertEqual(resolve_token(" abc "), "abc")
        finally:
            del os.environ["BRAPI_TOKEN"]


class RankingPriceTests(unittest.TestCase):
    def setUp(self):
        self.r = rec_file([("AAAA3", 9, 1, "1.000.000"), ("EXAT3", 8, 2, "1.000.000"),
                           ("BAIX3", 8, 2, "1.000.000"), ("NOPR3", 8, 2, "1.000.000"),
                           ("LOWV3", 9, 1, "100")])
        self.prices = {"AAAA3": 12.5, "EXAT3": 4.0, "BAIX3": 3.99, "LOWV3": 1.0}

    def test_filter_boundary_and_missing(self):
        res = rank_records(self.r.records, ScoringConfig(min_price=4.0), self.prices)
        self.assertEqual({x.record.ticker for x in res.ranked}, {"AAAA3", "EXAT3"})   # 4,00 exato passa
        reasons = {e.record.ticker: e.reason for e in res.excluded_price}
        self.assertEqual(reasons, {"BAIX3": "Valor da ação abaixo do mínimo", "NOPR3": "Sem cotação"})
        self.assertEqual((res.price_low, res.price_missing), (1, 1))
        self.assertEqual(res.total, len(res.excluded_volume) + len(res.excluded_filters)
                         + len(res.excluded_price) + res.eligible)
        self.assertEqual([e.record.ticker for e in res.excluded_volume], ["LOWV3"])  # volume vem antes do preço

    def test_no_price_filter_keeps_old_behaviour(self):
        res = rank_records(self.r.records, ScoringConfig())
        self.assertEqual(res.eligible, 4)
        self.assertEqual(res.excluded_price, [])
        self.assertTrue(all(row["price"] is None for row in ranking_rows(res)))

    def test_price_column_in_rows_and_exports(self):
        res = rank_records(self.r.records, ScoringConfig(min_price=4.0), self.prices)
        rows = ranking_rows(res)
        self.assertEqual({r["ticker"]: r["price"] for r in rows}, {"AAAA3": 12.5, "EXAT3": 4.0})
        self.assertEqual(fmt_value("money", 12.5), "R$ 12,50")
        self.assertEqual(fmt_value("money", None), "—")
        self.assertIn("price", {c.key for c in RANKING_COLUMNS})
        ex = excluded_rows(res)
        self.assertEqual({e["ticker"]: e["reason"] for e in ex}["NOPR3"], "Sem cotação")
        d = tempfile.mkdtemp()
        export_csv(os.path.join(d, "a.csv"), RANKING_COLUMNS, rows)
        with open(os.path.join(d, "a.csv"), encoding="utf-8-sig") as fh:
            txt = fh.read()
        self.assertIn("Valor Ação (R$)", txt)
        self.assertIn(";12,50;", txt)
        export_xlsx(os.path.join(d, "a.xlsx"), [("Compra", RANKING_COLUMNS, rows)])
        from openpyxl import load_workbook
        ws = load_workbook(os.path.join(d, "a.xlsx"))["Compra"]
        self.assertEqual(ws["C2"].value, 12.5)                    # número real
        self.assertIn("R$", ws["C2"].number_format)

    def test_compare_with_price(self):
        a = rec_file([("AAAA3", 9, 1, "1.000.000"), ("BAIX3", 8, 2, "1.000.000"), ("NOPR3", 8, 2, "1.000.000"),
                      ("FEWO3", 2, 1, "1.000.000")])
        b = rec_file([("AAAA3", 12, 3, "1.000.000"), ("BAIX3", 9, 3, "1.000.000"), ("NOPR3", 9, 3, "1.000.000"),
                      ("FEWO3", 8, 4, "1.000.000")])
        prices = {"AAAA3": 10.0, "BAIX3": 2.0, "FEWO3": 9.0}
        c = compare(a, b, ScoringConfig(min_price=4.0), prices=prices)
        self.assertEqual([m.ticker for m in c.matches], ["AAAA3"])
        self.assertEqual(c.matches[0].price, 10.0)
        self.assertEqual((c.dropped_price, c.dropped_filters, c.dropped_volume), (2, 1, 0))
        self.assertEqual(c.n_both, c.dropped_volume + c.dropped_filters + c.dropped_price
                         + c.dropped_top_n + len(c.matches))
        self.assertEqual(comparison_rows(c.matches)[0]["price"], 10.0)
        labels = [k for k, _ in comparison_summary(c)]
        self.assertTrue(any("valor da ação" in k for k in labels))
        self.assertIn("price", {col.key for col in comparison_columns("2%", "1%")})
        c0 = compare(a, b, ScoringConfig())                       # sem filtro: resumo sem linha de preço
        self.assertEqual(c0.dropped_price, 0)
        self.assertFalse(any("valor da ação" in k for k, _ in comparison_summary(c0)))

    def test_negative_min_price_rejected(self):
        with self.assertRaises(ValueError):
            ScoringConfig(min_price=-1).validate()


class SettingsTests(unittest.TestCase):
    def test_roundtrip_and_missing(self):
        p = os.path.join(tempfile.mkdtemp(), "sub", "s.json")
        self.assertEqual(load_settings(p), {})
        self.assertTrue(save_settings({"min_price": "4,00", "use_price": True}, p))
        self.assertEqual(load_settings(p), {"min_price": "4,00", "use_price": True})
        with open(p, "w") as fh:
            fh.write("{quebrado")
        self.assertEqual(load_settings(p), {})


if __name__ == "__main__":
    unittest.main()
